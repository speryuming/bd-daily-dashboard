#!/usr/bin/env python3
"""Local issue tracker for Fuquan store rectification."""

from __future__ import annotations

import argparse
import cgi
import hashlib
import io
import json
import os
import re
import secrets
import sqlite3
import threading
import time
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from http.cookies import SimpleCookie
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
UPLOADS = DATA / "uploads"
DB_PATH = DATA / "rectification.sqlite3"
STATIC = ROOT / "static"
APP_VERSION = "2026.09.23.5"
MAX_UPLOAD = 30 * 1024 * 1024
STATUS = {"待反馈", "待复核", "未通过", "待核查", "通过", "例外待决策"}
TYPES = (
    "爆品团商品对标", "神抢手商品对标", "超抢手商品对标", "商品价格",
    "配送费", "起送价", "一口价", "活动商品数量", "营业时间", "店名",
    "满减活动", "其他",
)
LOGIN_ATTEMPTS: dict[str, list[float]] = {}
LOGIN_LOCK = threading.Lock()
LOGIN_WINDOW_SECONDS = 15 * 60
LOGIN_MAX_FAILURES = 8


def now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def today() -> str:
    return date.today().isoformat()


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def login_is_limited(client: str) -> bool:
    cutoff = time.time() - LOGIN_WINDOW_SECONDS
    with LOGIN_LOCK:
        attempts = [stamp for stamp in LOGIN_ATTEMPTS.get(client, []) if stamp >= cutoff]
        LOGIN_ATTEMPTS[client] = attempts
        return len(attempts) >= LOGIN_MAX_FAILURES


def record_login_failure(client: str) -> None:
    with LOGIN_LOCK:
        LOGIN_ATTEMPTS.setdefault(client, []).append(time.time())


def clear_login_failures(client: str) -> None:
    with LOGIN_LOCK:
        LOGIN_ATTEMPTS.pop(client, None)


def normalized(value: str) -> str:
    return re.sub(r"\s+", "", clean(value)).lower()


def parse_date(value, year: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = clean(value)
    if not s:
        return ""
    nums = re.findall(r"\d+", s)
    try:
        if len(nums) >= 3 and len(nums[0]) == 4:
            return date(int(nums[0]), int(nums[1]), int(nums[2])).isoformat()
        if len(nums) >= 2:
            return date(year, int(nums[0]), int(nums[1])).isoformat()
    except ValueError:
        return ""
    return ""


def split_issues(value: str) -> list[str]:
    text = clean(value).replace("\r", "\n")
    if not text:
        return []
    text = re.sub(r"(?m)(^|[\n，,;；])\s*\d+\s*[：:、，,.]\s*", "\n§", text)
    chunks = [x.strip(" \n,，;；") for x in text.split("§") if x.strip(" \n,，;；")]
    return chunks


def classify(text: str, sheet: str = "") -> str:
    t = clean(text)
    if "配送费" in t or "免配" in t or "配送高于" in t:
        return "配送费"
    if "起送" in t:
        return "起送价"
    if "神抢手" in t:
        return "神抢手商品对标"
    if "超抢手" in t or "超枪手" in t or "超枪手" in sheet:
        return "超抢手商品对标"
    if "爆品团" in t or "拼团" in sheet:
        return "爆品团商品对标"
    if "一口价" in t:
        return "一口价"
    if "价格" in t or "到手价" in t:
        return "商品价格"
    if "商品数量" in t or "产品数量" in t or "热销品" in t:
        return "活动商品数量"
    if "营业时间" in t or "未营业" in t:
        return "营业时间"
    if "店名" in t:
        return "店名"
    if "满减" in t:
        return "满减活动"
    return "其他"


def parse_workbook(path: Path, sheet_name: str, year: int) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError("找不到工作表：" + sheet_name)
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = [clean(x) for x in next(rows, [])]
        result = []
        if "店铺名称" in header and "问题描述" in header:
            columns = {name: header.index(name) for name in header if name}
            def cell(row, key):
                i = columns.get(key)
                return clean(row[i]) if i is not None and i < len(row) else ""
            for row_no, row in enumerate(rows, 2):
                store = cell(row, "店铺名称")
                if not store:
                    continue
                raw_date = cell(row, "日期") or cell(row, "时间")
                discovered = parse_date(raw_date, year)
                desc = cell(row, "问题描述")
                type_text = cell(row, "问题类型")
                type_parts = [x.strip() for x in re.split(r"[\n;；]+", type_text) if x.strip()]
                parts = type_parts + split_issues(desc)
                if not parts:
                    continue
                for item_index, part in enumerate(parts, 1):
                    result.append({
                        "source_row": row_no, "item_index": item_index,
                        "discovered_date": discovered, "raw_date": raw_date,
                        "raw_store_name": store, "issue_text": part,
                        "issue_type": classify(part, sheet_name),
                        "owner": cell(row, "责任BD") or cell(row, "BD"),
                        "source_feedback": cell(row, "是否整改") or cell(row, "整改状态"),
                        "source_review": cell(row, "运营复核") or cell(row, "复核"),
                        "product_name": "", "competitor_sales": "", "own_sales": "",
                        "competitor_price": "", "own_price": "",
                    })
        elif "商品名称" in header and "店铺名称" in header:
            columns = {name: header.index(name) for name in header if name}
            def cell(row, key):
                i = columns.get(key)
                return clean(row[i]) if i is not None and i < len(row) else ""
            last_store = last_date = ""
            for row_no, row in enumerate(rows, 2):
                last_store = cell(row, "店铺名称") or last_store
                last_date = cell(row, "时间") or cell(row, "日期") or last_date
                product = cell(row, "商品名称")
                if not product or not last_store:
                    continue
                result.append({
                    "source_row": row_no, "item_index": 1,
                    "discovered_date": parse_date(last_date, year), "raw_date": last_date,
                    "raw_store_name": last_store, "issue_text": product,
                    "issue_type": classify(product, sheet_name), "owner": cell(row, "BD"),
                    "source_feedback": cell(row, "是否整改") or cell(row, "bd是否整改"),
                    "source_review": "", "product_name": product,
                    "competitor_sales": cell(row, "竞对销量"),
                    "own_sales": cell(row, "ET销量"),
                    "competitor_price": cell(row, "竞对价格"),
                    "own_price": cell(row, "ET价格"),
                })
        else:
            raise ValueError("该工作表表头暂不支持；请选择综合问题表或商品对标表")
        return result
    finally:
        wb.close()


def connect() -> sqlite3.Connection:
    DATA.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH, timeout=10)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA busy_timeout=5000")
    return db


def init_db() -> None:
    with connect() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS issues (
          id INTEGER PRIMARY KEY, discovered_date TEXT NOT NULL,
          raw_store_name TEXT NOT NULL, issue_text TEXT NOT NULL,
          issue_type TEXT NOT NULL, owner TEXT NOT NULL DEFAULT '',
          status TEXT NOT NULL DEFAULT '待反馈', due_date TEXT NOT NULL DEFAULT '',
          review_due_date TEXT NOT NULL DEFAULT '', latest_seen TEXT NOT NULL DEFAULT '',
          product_name TEXT NOT NULL DEFAULT '', competitor_sales TEXT NOT NULL DEFAULT '',
          own_sales TEXT NOT NULL DEFAULT '', competitor_price TEXT NOT NULL DEFAULT '',
          own_price TEXT NOT NULL DEFAULT '', source_sheet TEXT NOT NULL DEFAULT '',
          source_row INTEGER NOT NULL DEFAULT 0, source_feedback TEXT NOT NULL DEFAULT '',
          source_review TEXT NOT NULL DEFAULT '', mapping_id INTEGER,
          created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
          FOREIGN KEY(mapping_id) REFERENCES mappings(id)
        );
        CREATE TABLE IF NOT EXISTS events (
          id INTEGER PRIMARY KEY, issue_id INTEGER NOT NULL, kind TEXT NOT NULL,
          actor TEXT NOT NULL DEFAULT '', action TEXT NOT NULL DEFAULT '',
          note TEXT NOT NULL DEFAULT '', evidence TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL, FOREIGN KEY(issue_id) REFERENCES issues(id)
        );
        CREATE TABLE IF NOT EXISTS imports (
          file_sha TEXT NOT NULL, sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
          item_index INTEGER NOT NULL, issue_id INTEGER NOT NULL,
          PRIMARY KEY(file_sha,sheet,source_row,item_index),
          FOREIGN KEY(issue_id) REFERENCES issues(id)
        );
        CREATE TABLE IF NOT EXISTS mappings (
          id INTEGER PRIMARY KEY, raw_store_name TEXT NOT NULL,
          own_merchant_id TEXT NOT NULL, own_store_name TEXT NOT NULL,
          competitor_store_name TEXT NOT NULL, competitor_url TEXT NOT NULL DEFAULT '',
          evidence TEXT NOT NULL DEFAULT '', confirmed_by TEXT NOT NULL,
          active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS managers (
          id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE,
          salt TEXT NOT NULL, password_hash TEXT NOT NULL,
          created_by TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS manager_sessions (
          token_hash TEXT PRIMARY KEY, manager_id INTEGER NOT NULL,
          expires_at INTEGER NOT NULL,
          FOREIGN KEY(manager_id) REFERENCES managers(id)
        );
        CREATE TABLE IF NOT EXISTS users (
          id INTEGER PRIMARY KEY, username TEXT NOT NULL UNIQUE,
          display_name TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('manager','staff')),
          salt TEXT NOT NULL, password_hash TEXT NOT NULL,
          active INTEGER NOT NULL DEFAULT 1,
          created_by TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS sessions (
          token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL,
          expires_at INTEGER NOT NULL,
          FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE INDEX IF NOT EXISTS idx_issues_status ON issues(status);
        CREATE INDEX IF NOT EXISTS idx_issues_store ON issues(raw_store_name);
        """)
        # Preserve manager accounts created by the single-user version.
        db.execute("""INSERT OR IGNORE INTO users
            (username,display_name,role,salt,password_hash,active,created_by,created_at)
            SELECT username,username,'manager',salt,password_hash,1,created_by,created_at FROM managers""")
        if db.execute("SELECT COUNT(*) FROM users").fetchone()[0] == 0:
            admin_username = clean(os.environ.get("RECTIFICATION_ADMIN_USERNAME"))
            admin_password = os.environ.get("RECTIFICATION_ADMIN_PASSWORD", "")
            admin_name = clean(os.environ.get("RECTIFICATION_ADMIN_NAME")) or admin_username
            if admin_username and admin_password:
                create_user(db, admin_username, admin_password, admin_name, "manager", "云端首次部署")


def password_hash(password: str, salt: bytes) -> str:
    return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 300_000).hex()


def create_user(db: sqlite3.Connection, username: str, password: str, display_name: str,
                role: str, created_by: str) -> None:
    username = clean(username)
    display_name = clean(display_name) or username
    if len(username) < 2 or len(username) > 40:
        raise ValueError("账号需为2至40个字符")
    if len(password) < 10:
        raise ValueError("密码至少10个字符")
    if role not in {"manager", "staff"}:
        raise ValueError("账号角色无效")
    salt = secrets.token_bytes(16)
    try:
        db.execute("""INSERT INTO users
            (username,display_name,role,salt,password_hash,created_by,created_at)
            VALUES (?,?,?,?,?,?,?)""",
            (username, display_name, role, salt.hex(), password_hash(password, salt), created_by, now()))
    except sqlite3.IntegrityError:
        raise ValueError("该账号已存在")


def create_manager(db: sqlite3.Connection, username: str, password: str, created_by: str) -> None:
    """Compatibility helper used by tests and older setup code."""
    create_user(db, username, password, username, "manager", created_by)


def user_from_cookie(db: sqlite3.Connection, cookie_header: str):
    cookies = SimpleCookie()
    try:
        cookies.load(cookie_header)
    except Exception:
        return None
    token = cookies.get("session")
    if not token:
        return None
    digest = hashlib.sha256(token.value.encode("utf-8")).hexdigest()
    return db.execute("""SELECT u.id,u.username,u.display_name,u.role FROM sessions s
        JOIN users u ON u.id=s.user_id
        WHERE s.token_hash=? AND s.expires_at>? AND u.active=1""",
        (digest, int(time.time()))).fetchone()


def require_user(db: sqlite3.Connection, cookie_header: str, role: str | None = None):
    user = user_from_cookie(db, cookie_header)
    if not user:
        raise PermissionError("请先登录")
    if role and user["role"] != role:
        raise PermissionError("此操作仅限管理人员")
    return user


def sheet_names(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [x for x in wb.sheetnames if not x.startswith("WpsReserved")]
    finally:
        wb.close()


def import_rows(db: sqlite3.Connection, rows: list[dict], sha: str, sheet: str) -> dict:
    counts = {"created": 0, "reused": 0, "already_imported": 0, "needs_date": 0}
    stamp = now()
    for row in rows:
        if not row["discovered_date"]:
            counts["needs_date"] += 1
            continue
        key = (sha, sheet, row["source_row"], row["item_index"])
        if db.execute("SELECT 1 FROM imports WHERE file_sha=? AND sheet=? AND source_row=? AND item_index=?", key).fetchone():
            counts["already_imported"] += 1
            continue
        match = db.execute("""SELECT id FROM issues WHERE raw_store_name=? AND issue_type=?
            AND issue_text=? AND status NOT IN ('通过') ORDER BY id DESC LIMIT 1""",
            (row["raw_store_name"], row["issue_type"], row["issue_text"])).fetchone()
        if match:
            issue_id = match["id"]
            db.execute("UPDATE issues SET latest_seen=?,updated_at=? WHERE id=?",
                       (row["discovered_date"], stamp, issue_id))
            counts["reused"] += 1
        else:
            cur = db.execute("""INSERT INTO issues
                (discovered_date,raw_store_name,issue_text,issue_type,owner,latest_seen,
                 product_name,competitor_sales,own_sales,competitor_price,own_price,
                 source_sheet,source_row,source_feedback,source_review,created_at,updated_at)
                 VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (row["discovered_date"], row["raw_store_name"], row["issue_text"],
                 row["issue_type"], row["owner"], row["discovered_date"],
                 row["product_name"], row["competitor_sales"], row["own_sales"],
                 row["competitor_price"], row["own_price"], sheet, row["source_row"],
                 row["source_feedback"], row["source_review"], stamp, stamp))
            issue_id = cur.lastrowid
            counts["created"] += 1
        db.execute("INSERT INTO imports VALUES (?,?,?,?,?)", (*key, issue_id))
        date_note = "（原表日期空白，按导入时指定日期补入）" if row.get("date_assumed") else ""
        db.execute("INSERT INTO events(issue_id,kind,actor,note,created_at) VALUES (?,?,?,?,?)",
                   (issue_id, "导入", "系统", f"{sheet} 第{row['source_row']}行：{row['issue_text']}{date_note}", stamp))
    return counts


def issue_dict(db: sqlite3.Connection, row: sqlite3.Row, events=False) -> dict:
    item = dict(row)
    if events:
        item["events"] = [dict(x) for x in db.execute(
            "SELECT * FROM events WHERE issue_id=? ORDER BY id DESC", (row["id"],))]
    return item


def create_manual_issue(db: sqlite3.Connection, data: dict, actor: str) -> int:
    discovered_date = clean(data.get("discovered_date"))
    store = clean(data.get("raw_store_name"))
    issue_type = clean(data.get("issue_type"))
    issue_text = clean(data.get("issue_text"))
    owner = clean(data.get("owner"))
    due_date = clean(data.get("due_date"))
    if not discovered_date or not store or not issue_text:
        raise ValueError("发现日期、店铺名称和问题描述必填")
    date.fromisoformat(discovered_date)
    if due_date:
        date.fromisoformat(due_date)
    if issue_type not in TYPES:
        raise ValueError("问题类型无效")
    if len(store) > 200 or len(issue_text) > 3000:
        raise ValueError("店铺名称或问题描述过长")
    if owner:
        owner_user = db.execute("""SELECT display_name FROM users
            WHERE active=1 AND (display_name=? OR username=?)""", (owner, owner)).fetchone()
        if not owner_user:
            raise ValueError("负责人必须选择一个有效账号")
        owner = owner_user["display_name"]
    stamp = now()
    cur = db.execute("""INSERT INTO issues
        (discovered_date,raw_store_name,issue_text,issue_type,owner,due_date,latest_seen,
         product_name,competitor_sales,own_sales,competitor_price,own_price,
         source_sheet,source_row,created_at,updated_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (discovered_date, store, issue_text, issue_type, owner, due_date, discovered_date,
         clean(data.get("product_name")), clean(data.get("competitor_sales")),
         clean(data.get("own_sales")), clean(data.get("competitor_price")),
         clean(data.get("own_price")), "人工新增", 0, stamp, stamp))
    issue_id = cur.lastrowid
    note = f"{issue_type}：{issue_text}"
    if owner:
        note += f"；负责人：{owner}"
    db.execute("INSERT INTO events(issue_id,kind,actor,note,created_at) VALUES (?,?,?,?,?)",
               (issue_id, "手动新增", actor, note, stamp))
    return issue_id


class Handler(BaseHTTPRequestHandler):
    server_version = "FuquanRectification/1.0"

    def log_message(self, fmt, *args):
        print("[%s] %s" % (self.log_date_time_string(), fmt % args))

    def respond(self, payload, status=200, cookie=None):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        if cookie:
            self.send_header("Set-Cookie", cookie)
        self.end_headers()
        self.wfile.write(body)

    def fail(self, message, status=400):
        self.respond({"error": message}, status)

    def client_ip(self):
        direct = self.client_address[0]
        if direct in {"127.0.0.1", "::1"}:
            return clean(self.headers.get("CF-Connecting-IP")) or direct
        return direct

    def session_cookie(self, token: str, max_age: int) -> str:
        secure = "; Secure" if self.headers.get("X-Forwarded-Proto", "").lower() == "https" else ""
        return f"session={token}; HttpOnly; SameSite=Strict; Path=/; Max-Age={max_age}{secure}"

    def read_json(self):
        size = int(self.headers.get("Content-Length", "0"))
        if size <= 0 or size > 2_000_000:
            raise ValueError("请求内容为空或过大")
        return json.loads(self.rfile.read(size))

    def authorize(self, db, role=None):
        user = user_from_cookie(db, self.headers.get("Cookie", ""))
        if not user:
            self.fail("请先登录", 401)
            return None
        if role and user["role"] != role:
            self.fail("此操作仅限管理人员", 403)
            return None
        return user

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/version":
            self.respond({"version": APP_VERSION})
            return
        if parsed.path == "/api/auth":
            with connect() as db:
                user = user_from_cookie(db, self.headers.get("Cookie", ""))
                count = db.execute("SELECT COUNT(*) FROM users").fetchone()[0]
                self.respond({"user": dict(user) if user else None,
                              "manager": dict(user) if user and user["role"] == "manager" else None,
                              "setup_needed": count == 0})
            return
        if parsed.path == "/api/users":
            with connect() as db:
                if not self.authorize(db, "manager"): return
                self.respond({"items": [dict(x) for x in db.execute(
                    "SELECT id,username,display_name,role,active FROM users ORDER BY role,display_name")]})
            return
        if parsed.path == "/api/issues":
            qs = parse_qs(parsed.query)
            status = qs.get("status", [""])[0]
            q = qs.get("q", [""])[0]
            with connect() as db:
                if not self.authorize(db): return
                where, args = [], []
                if status and status != "全部":
                    where.append("i.status=?"); args.append(status)
                if q:
                    where.append("(i.raw_store_name LIKE ? OR i.issue_text LIKE ? OR i.owner LIKE ?)")
                    args.extend([f"%{q}%"] * 3)
                sql = """SELECT i.*,m.own_merchant_id,m.own_store_name,m.competitor_store_name,
                    m.competitor_url FROM issues i LEFT JOIN mappings m ON i.mapping_id=m.id"""
                if where:
                    sql += " WHERE " + " AND ".join(where)
                sql += " ORDER BY CASE WHEN i.status='通过' THEN 1 ELSE 0 END, i.discovered_date DESC, i.id DESC LIMIT 1000"
                self.respond({"items": [dict(x) for x in db.execute(sql, args)]})
            return
        if parsed.path == "/api/summary":
            with connect() as db:
                if not self.authorize(db): return
                counts = {r["status"]: r["n"] for r in db.execute("SELECT status,COUNT(*) n FROM issues GROUP BY status")}
                self.respond({"counts": counts, "total": sum(counts.values()), "today": today(),
                              "unmapped": db.execute("SELECT COUNT(*) FROM issues WHERE mapping_id IS NULL").fetchone()[0]})
            return
        if parsed.path == "/api/mappings":
            with connect() as db:
                if not self.authorize(db, "manager"): return
                self.respond({"items": [dict(x) for x in db.execute("SELECT * FROM mappings WHERE active=1 ORDER BY id DESC")]})
            return
        m = re.fullmatch(r"/api/issues/(\d+)", parsed.path)
        if m:
            with connect() as db:
                if not self.authorize(db): return
                row = db.execute("""SELECT i.*,m.own_merchant_id,m.own_store_name,m.competitor_store_name,
                    m.competitor_url FROM issues i LEFT JOIN mappings m ON i.mapping_id=m.id WHERE i.id=?""", (int(m[1]),)).fetchone()
                if not row:
                    self.fail("问题不存在", 404)
                else:
                    self.respond(issue_dict(db, row, True))
            return
        path = "index.html" if parsed.path == "/" else parsed.path.lstrip("/")
        target = (STATIC / path).resolve()
        if not target.is_file() or STATIC.resolve() not in target.parents:
            self.send_error(404); return
        mime = "text/html" if target.suffix == ".html" else "text/css" if target.suffix == ".css" else "text/javascript"
        content = target.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", mime + "; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Content-Security-Policy", "default-src 'self'; style-src 'self'; script-src 'self'; connect-src 'self'; form-action 'self'; frame-ancestors 'none'")
        self.end_headers()
        self.wfile.write(content)

    def do_POST(self):
        if self.headers.get("X-Rectification-App") != "1":
            self.fail("缺少本地系统请求标记", 403); return
        try:
            self.post_route()
        except PermissionError as exc:
            self.fail(str(exc), 403)
        except (ValueError, json.JSONDecodeError) as exc:
            self.fail(str(exc))
        except Exception as exc:
            self.fail("操作失败：" + str(exc), 500)

    def post_route(self):
        path = urlparse(self.path).path
        if path in {"/api/setup-manager", "/api/login", "/api/logout", "/api/users", "/api/managers"}:
            data = self.read_json()
            with connect() as db:
                if path == "/api/setup-manager":
                    db.execute("BEGIN IMMEDIATE")
                    if db.execute("SELECT COUNT(*) FROM users").fetchone()[0]:
                        raise PermissionError("管理人员已设置，请登录")
                    create_user(db, clean(data.get("username")), clean(data.get("password")),
                                clean(data.get("display_name")), "manager", "首次设置")
                    self.respond({"ok": True})
                elif path == "/api/login":
                    username = clean(data.get("username"))
                    password = clean(data.get("password"))
                    client = self.client_ip()
                    if login_is_limited(client):
                        raise PermissionError("登录尝试次数过多，请15分钟后再试")
                    row = db.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
                    if not row or not secrets.compare_digest(password_hash(password, bytes.fromhex(row["salt"])), row["password_hash"]):
                        record_login_failure(client)
                        raise PermissionError("账号或密码错误")
                    clear_login_failures(client)
                    token = secrets.token_urlsafe(32)
                    digest = hashlib.sha256(token.encode("utf-8")).hexdigest()
                    db.execute("INSERT INTO sessions(token_hash,user_id,expires_at) VALUES (?,?,?)",
                               (digest, row["id"], int(time.time()) + 12 * 3600))
                    self.respond({"ok": True, "username": row["username"], "role": row["role"]},
                                 cookie=self.session_cookie(token, 43200))
                elif path == "/api/logout":
                    cookies = SimpleCookie()
                    cookies.load(self.headers.get("Cookie", ""))
                    if cookies.get("session"):
                        digest = hashlib.sha256(cookies["session"].value.encode("utf-8")).hexdigest()
                        db.execute("DELETE FROM sessions WHERE token_hash=?", (digest,))
                    self.respond({"ok": True}, cookie=self.session_cookie("", 0))
                else:
                    manager = require_user(db, self.headers.get("Cookie", ""), "manager")
                    create_user(db, clean(data.get("username")), clean(data.get("password")),
                                clean(data.get("display_name")), clean(data.get("role")) or "staff",
                                manager["display_name"])
                    self.respond({"ok": True})
            return
        if path == "/api/issues/manual":
            data = self.read_json()
            with connect() as db:
                manager = require_user(db, self.headers.get("Cookie", ""), "manager")
                issue_id = create_manual_issue(db, data, manager["display_name"])
            self.respond({"ok": True, "id": issue_id})
            return
        if path == "/api/upload":
            with connect() as db:
                require_user(db, self.headers.get("Cookie", ""), "manager")
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > MAX_UPLOAD:
                raise ValueError("文件为空或超过30MB")
            if not self.headers.get("Content-Type", "").startswith("multipart/form-data"):
                raise ValueError("请上传 xlsx 文件")
            form = cgi.FieldStorage(fp=io.BytesIO(self.rfile.read(length)),
                headers=self.headers, environ={"REQUEST_METHOD": "POST", "CONTENT_TYPE": self.headers["Content-Type"]})
            if "file" not in form:
                raise ValueError("未选择文件")
            upload = form["file"]
            filename = clean(upload.filename)
            if not filename.lower().endswith(".xlsx"):
                raise ValueError("只支持 xlsx 文件")
            content = upload.file.read()
            if not content.startswith(b"PK"):
                raise ValueError("文件不是有效的 xlsx")
            sha = hashlib.sha256(content).hexdigest()
            UPLOADS.mkdir(parents=True, exist_ok=True)
            dest = UPLOADS / (sha + ".xlsx")
            if not dest.exists():
                dest.write_bytes(content)
            self.respond({"token": sha, "filename": filename, "sheets": sheet_names(dest)})
            return
        if path == "/api/import-preview" or path == "/api/import-commit":
            with connect() as db:
                require_user(db, self.headers.get("Cookie", ""), "manager")
            data = self.read_json()
            token = clean(data.get("token"))
            sheet = clean(data.get("sheet"))
            year = int(data.get("year", date.today().year))
            if not re.fullmatch(r"[0-9a-f]{64}", token) or not 2020 <= year <= 2100:
                raise ValueError("文件标识或年份无效")
            source = UPLOADS / (token + ".xlsx")
            if not source.exists():
                raise ValueError("上传文件不存在，请重新上传")
            rows = parse_workbook(source, sheet, year)
            fallback_date = clean(data.get("fallback_date"))
            if fallback_date:
                date.fromisoformat(fallback_date)
                for row in rows:
                    if not row["discovered_date"]:
                        row["discovered_date"] = fallback_date
                        row["date_assumed"] = True
            if path.endswith("preview"):
                self.respond({"rows": rows[:200], "total": len(rows),
                              "needs_date": sum(not x["discovered_date"] for x in rows)})
            else:
                with connect() as db:
                    result = import_rows(db, rows, token, sheet)
                self.respond({"result": result})
            return
        m = re.fullmatch(r"/api/issues/(\d+)/(assign|feedback|review|map|link-mapping)", path)
        if not m:
            self.fail("接口不存在", 404); return
        issue_id, operation = int(m[1]), m[2]
        data = self.read_json()
        with connect() as db:
            user = require_user(db, self.headers.get("Cookie", ""))
            issue = db.execute("SELECT * FROM issues WHERE id=?", (issue_id,)).fetchone()
            if not issue:
                self.fail("问题不存在", 404); return
            stamp = now()
            if operation == "assign":
                if user["role"] != "manager":
                    raise PermissionError("此操作仅限管理人员")
                owner = clean(data.get("owner"))
                due = clean(data.get("due_date"))
                if not owner:
                    raise ValueError("请填写负责人")
                owner_user = db.execute("""SELECT display_name FROM users
                    WHERE active=1 AND (display_name=? OR username=?)""", (owner, owner)).fetchone()
                if not owner_user:
                    raise ValueError("负责人必须选择一个有效账号")
                owner = owner_user["display_name"]
                if due:
                    date.fromisoformat(due)
                db.execute("UPDATE issues SET owner=?,due_date=?,updated_at=? WHERE id=?", (owner, due, stamp, issue_id))
                db.execute("INSERT INTO events(issue_id,kind,actor,note,created_at) VALUES (?,?,?,?,?)",
                           (issue_id, "分派", user["display_name"], f"负责人：{owner}；截止：{due or '未设'}", stamp))
            elif operation == "feedback":
                if user["role"] != "manager" and issue["owner"] not in {user["username"], user["display_name"]}:
                    raise PermissionError("只能反馈分配给自己的问题")
                actor = user["display_name"]
                action = clean(data.get("action"))
                note = clean(data.get("note"))
                if not action or not note:
                    raise ValueError("处理动作和说明必填")
                due = (date.today() + timedelta(days=1)).isoformat()
                db.execute("UPDATE issues SET status='待复核',review_due_date=?,updated_at=? WHERE id=?",
                           (due, stamp, issue_id))
                db.execute("INSERT INTO events(issue_id,kind,actor,action,note,evidence,created_at) VALUES (?,?,?,?,?,?,?)",
                           (issue_id, "反馈", actor, action, note, clean(data.get("evidence")), stamp))
            elif operation == "review":
                if user["role"] != "manager":
                    raise PermissionError("只有管理人员登录后才能复核")
                actor = user["display_name"]
                result = clean(data.get("result"))
                note = clean(data.get("note"))
                if not note or result not in {"通过", "未通过", "待核查", "例外待决策"}:
                    raise ValueError("复核结论和说明必填")
                if issue["status"] not in {"待复核", "未通过", "待核查", "例外待决策"}:
                    raise ValueError("请先提交整改反馈，再进行复核")
                if result == "通过" and issue["review_due_date"] and today() < issue["review_due_date"]:
                    raise ValueError("计划复核日尚未到，不能提前判定通过")
                status = "待反馈" if result == "未通过" else result
                db.execute("UPDATE issues SET status=?,updated_at=? WHERE id=?", (status, stamp, issue_id))
                db.execute("INSERT INTO events(issue_id,kind,actor,action,note,evidence,created_at) VALUES (?,?,?,?,?,?,?)",
                           (issue_id, "复核", actor, result, note, clean(data.get("evidence")), stamp))
            elif operation == "map":
                if user["role"] != "manager":
                    raise PermissionError("此操作仅限管理人员")
                fields = {k: clean(data.get(k)) for k in ("own_merchant_id", "own_store_name", "competitor_store_name", "competitor_url", "evidence", "confirmed_by")}
                if not fields["own_merchant_id"] or not fields["competitor_store_name"] or not fields["confirmed_by"]:
                    raise ValueError("我方商户ID、美团店铺和确认人必填")
                cur = db.execute("""INSERT INTO mappings(raw_store_name,own_merchant_id,own_store_name,
                    competitor_store_name,competitor_url,evidence,confirmed_by,created_at)
                    VALUES (?,?,?,?,?,?,?,?)""", (issue["raw_store_name"], *(fields[k] for k in fields), stamp))
                db.execute("UPDATE issues SET mapping_id=?,updated_at=? WHERE id=?", (cur.lastrowid, stamp, issue_id))
                db.execute("INSERT INTO events(issue_id,kind,actor,note,created_at) VALUES (?,?,?,?,?)",
                           (issue_id, "人工映射", fields["confirmed_by"], f"我方ID {fields['own_merchant_id']} ↔ 美团 {fields['competitor_store_name']}", stamp))
            else:
                if user["role"] != "manager":
                    raise PermissionError("此操作仅限管理人员")
                mapping_id = int(data.get("mapping_id", 0))
                actor = user["display_name"]
                mapping = db.execute("SELECT * FROM mappings WHERE id=? AND active=1", (mapping_id,)).fetchone()
                if not mapping or not actor:
                    raise ValueError("请选择人工确认过的映射并填写确认人")
                db.execute("UPDATE issues SET mapping_id=?,updated_at=? WHERE id=?", (mapping_id, stamp, issue_id))
                db.execute("INSERT INTO events(issue_id,kind,actor,note,created_at) VALUES (?,?,?,?,?)",
                           (issue_id, "人工映射", actor, f"沿用已确认映射 #{mapping_id}", stamp))
        self.respond({"ok": True})


def main():
    global DATA, UPLOADS, DB_PATH
    parser = argparse.ArgumentParser(description="福泉整改跟进系统")
    parser.add_argument("--host", default=os.environ.get("HOST", "127.0.0.1"))
    parser.add_argument("--port", type=int, default=int(os.environ.get("PORT", "8765")))
    parser.add_argument("--data-dir", type=Path,
                        default=Path(os.environ.get("RECTIFICATION_DATA_DIR", str(DATA))))
    args = parser.parse_args()
    DATA = args.data_dir.resolve()
    UPLOADS = DATA / "uploads"
    DB_PATH = DATA / "rectification.sqlite3"
    init_db()
    server = ThreadingHTTPServer((args.host, args.port), Handler)
    server.daemon_threads = True
    print(f"打开 http://127.0.0.1:{args.port}/", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
